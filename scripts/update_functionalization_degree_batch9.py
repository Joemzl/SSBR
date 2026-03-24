#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
第九批官能化程度数据更新 - 基于Zotero文献全文提取
更新样本: SSBR-090, SSBR-093, SSBR-097, SSBR-104, SSBR-107, SSBR-108
"""

import openpyxl
from pathlib import Path

# 官能化程度数据 - 从Zotero文献中提取
# 格式: (样本ID, 数值, 单位, 来源说明)
FUNCTIONALIZATION_DATA = [
    # SSBR-090: UPy-NCO 官能化
    # DOI: 10.1016/j.polymer.2024.128729
    # Table 1: UPy-NCO content = 1-3 mol% (based on Bd units)
    ("SSBR-090", "1-3", "mol% (vs Bd units)", "Table 1: UPy-NCO content 1-3 mol% based on Bd units"),
    
    # SSBR-093: 硫醇官能化GO (MPTMS修饰)
    # DOI: 10.1016/j.compscitech.2022.109907
    # Table 1: MGO = 2 phr
    ("SSBR-093", "2", "phr", "Table 1: MGO 2 phr (MPTMS-modified GO)"),
    
    # SSBR-097: 烃类树脂
    # DOI: 10.5254/RCT-D-23-00043
    # Table III: 树脂用量 = 20+10 = 30 phr (first pass + second pass)
    ("SSBR-097", "30", "phr", "Table III: resin 20+10=30 phr (DCPD/PMR types)"),
    
    # SSBR-104: 环氧乙烷链端封端
    # DOI: 10.1002/marc.202200366
    # 使用过量EO进行链端封端，制备双羟基封端SSBR (Mn~4000)
    ("SSBR-104", "双官能度", "(HT-SSBR)", "excess EO for chain-end hydroxyl termination, Mn~4000"),
    
    # SSBR-107: C3N4/SiO2 纳米杂化
    # DOI: 10.1002/pc.26994
    # Table 1: C3N4/SiO2 = 70 phr, 最佳C3N4含量 = 5 wt% (in hybrid)
    ("SSBR-107", "5", "wt% (C3N4 in hybrid)", "Table 1: 5% C3N4/SiO2 (70 phr total hybrid, 5 wt% C3N4)"),
    
    # SSBR-108: rGO-g-SiO2 接枝杂化
    # DOI: 10.1002/pc.26830
    # 最佳配方: 1.5 wt% rGO
    ("SSBR-108", "1.5", "wt% (rGO)", "optimal 1.5 wt% rGO in rGO-g-SiO2 hybrid"),
]

def update_excel():
    """更新Excel中的官能化程度数据"""
    excel_path = Path("d:/SSBR/dataset/数据.xlsx")
    
    # 加载工作簿
    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active
    
    # 获取列索引 (假设第一行是表头)
    headers = {cell.value: cell.column for cell in ws[1]}
    
    # 需要更新的列
    sample_col = headers.get("样本ID")
    degree_value_col = headers.get("官能化程度_原始数值")
    degree_unit_col = headers.get("官能化程度_原始单位")
    
    if not all([sample_col, degree_value_col, degree_unit_col]):
        print("错误: 找不到必要的列")
        print(f"Available headers: {list(headers.keys())}")
        return False
    
    # 创建样本ID到行号的映射
    sample_rows = {}
    for row in range(2, ws.max_row + 1):
        sample_id = ws.cell(row=row, column=sample_col).value
        if sample_id:
            sample_rows[sample_id] = row
    
    # 更新数据
    updated = []
    skipped = []
    
    for sample_id, value, unit, source in FUNCTIONALIZATION_DATA:
        if sample_id in sample_rows:
            row = sample_rows[sample_id]
            
            # 检查当前值
            current_value = ws.cell(row=row, column=degree_value_col).value
            
            # 更新数据
            ws.cell(row=row, column=degree_value_col).value = value
            ws.cell(row=row, column=degree_unit_col).value = unit
            
            updated.append((sample_id, value, unit, current_value))
            print(f"[OK] {sample_id}: {value} {unit}")
            print(f"  来源: {source}")
            if current_value and current_value != "待补充":
                print(f"  (原值: {current_value})")
        else:
            skipped.append(sample_id)
            print(f"[SKIP] {sample_id}: 未找到")
    
    # 保存
    wb.save(excel_path)
    
    print(f"\n=== 更新完成 ===")
    print(f"成功更新: {len(updated)} 个样本")
    print(f"跳过: {len(skipped)} 个样本")
    
    return True

if __name__ == "__main__":
    update_excel()
